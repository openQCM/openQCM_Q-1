"""
Build the Word edition of the openQCM Q-1 software user manual from its
Markdown master.

Run from the repository root:

    python docs/tools/build_manual_docx.py

The Markdown file is the source of truth. The existing .docx next to it is
used as the template: its styles, page setup, header/footer and cover-page
layout are kept, and its body is rebuilt from the Markdown. The result
replaces the .docx in place, so the next run uses it as its template.

Contents-page numbers are PAGEREF fields. They show "?" until Word updates
the fields, which export_manual_pdf.applescript does before saving the PDF.

Supported Markdown, i.e. what the manual uses:
- headings: "##" becomes Heading 1 and "###" Heading 2; on the cover, "#"
  is the title;
- paragraphs (single newlines become line breaks), **bold** and *italic*;
- "1." numbered and "-" bulleted items, written as plain hanging-indent
  paragraphs, as in the original Word edition;
- pipe tables, fenced code blocks, standalone [text](url) link paragraphs;
- images ![alt](path) and <img src=... width=px> (px at 96 dpi);
- <a id="..."></a> anchors, which become bookmarks on the next heading.
"""
import copy
import datetime
import os
import re
import sys

import docx
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MANUAL_DIR = os.path.join(REPO_ROOT, 'docs', 'user-manual')
BASENAME = 'openQCM_Q-1_Software_User_Manual_v3.0'
MD_PATH = os.path.join(MANUAL_DIR, BASENAME + '.md')
DOCX_PATH = os.path.join(MANUAL_DIR, BASENAME + '.docx')

TEXT_WIDTH_TWIPS = 9994          # 8.5 in page minus 2 x 1123 twips margins
TEXT_WIDTH_IN = TEXT_WIDTH_TWIPS / 1440
SCHEMATIC_WIDTH_IN = 6.9         # width used for the schematics in Edition 1.2
# GUI screenshots share one physical scale: the 1600 px main window spans
# the text width, so GUI text has the same size in every figure.
GUI_PX_PER_IN = 1600 / TEXT_WIDTH_IN
# Small detail crops (menus, single groups) at that scale would be a couple
# of centimetres wide; print them larger so their text stays readable.
DETAIL_MAX_PX = 900
DETAIL_PX_PER_IN = 200

LINK_COLOR = '214F6F'
CODE_FONT = 'Courier New'
TABLE_HEAD_FILL = '263F50'
TABLE_BAND_FILL = 'F3F6F8'
TABLE_BORDER = 'D9D9D9'

# Heading 2 sections that start on a new page (long firmware procedures).
H2_PAGE_BREAKS = {
    'Firmware update workflow',
    'Program the firmware with the external loader',
    'Verify the update and recover safely',
}
# From this chapter on, plain prose paragraphs use the Small style, except
# under the sub-headings listed below (same as the original Word edition).
SMALL_FROM_H1 = '24 References and provenance'
NORMAL_H2_IN_SMALL_ZONE = {'Document revision', 'Support and software citation'}


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------

def parse_blocks(text):
    """Split the Markdown into a flat list of (kind, payload) blocks."""
    lines = text.split('\n')
    blocks, para, i = [], [], 0

    def flush():
        if para:
            blocks.append(('para', '\n'.join(para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith('```'):
            flush()
            code = []
            i += 1
            while not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            blocks.append(('code', code))
        elif not s:
            flush()
        elif m := re.match(r'<a id="([^"]+)"></a>$', s):
            flush()
            blocks.append(('anchor', m.group(1)))
        elif m := re.match(r'(#{1,3}) (.+)$', s):
            flush()
            blocks.append(('h%d' % len(m.group(1)), m.group(2).strip()))
        elif m := re.match(r'!\[([^\]]*)\]\(([^)]+)\)$', s):
            flush()
            blocks.append(('img', {'alt': m.group(1), 'src': m.group(2), 'width_px': None}))
        elif s.startswith('<img '):
            flush()
            attrs = dict(re.findall(r'(\w+)="([^"]*)"', s))
            blocks.append(('img', {'alt': attrs.get('alt', ''), 'src': attrs['src'],
                                   'width_px': int(attrs['width']) if 'width' in attrs else None}))
        elif s.startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = lines[i].strip()
                if not re.match(r'^\|[\s\-:|]+\|$', row):
                    cells = re.split(r'(?<!\\)\|', row.strip('|'))
                    rows.append([c.strip().replace('\\|', '|') for c in cells])
                i += 1
            blocks.append(('table', rows))
            continue
        elif m := re.match(r'(\d+)\. +(.*)$', s):
            flush()
            blocks.append(('num', (m.group(1), m.group(2))))
        elif m := re.match(r'- +(.*)$', s):
            flush()
            blocks.append(('bullet', m.group(1)))
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


INLINE = re.compile(r'(\*\*.+?\*\*(?!\*)|\*\*.+?\*\*\*|\*[^*\s][^*]*\*)')


def inline_runs(text):
    """Split text into (segment, bold, italic) runs. Handles **bold**,
    *italic* and *italic* nested at the end of bold (**a *b***)."""
    text = text.replace('\\[', '[').replace('\\]', ']')
    runs = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**'):
            inner = part[2:-2]
            for sub in re.split(r'(\*[^*]+\*)', inner):
                if sub.startswith('*') and sub.endswith('*') and len(sub) > 1:
                    runs.append((sub[1:-1], True, True))
                elif sub:
                    runs.append((sub, True, False))
        elif part.startswith('*'):
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------

class Builder:
    def __init__(self, template):
        self.doc = docx.Document(template)
        self.body = self.doc.element.body
        self.sect = self.body.find(qn('w:sectPr'))
        self.bookmark_id = 100

    # --- low-level -------------------------------------------------------
    def new_par(self, style=None):
        p = self.doc.add_paragraph()
        self.body.remove(p._p)
        self.sect.addprevious(p._p)
        if style:
            p.style = self.doc.styles[style]
        return p

    def add_text(self, p, text, size=None, color=None, bold=False):
        for line_no, line in enumerate(text.split('\n')):
            if line_no:
                p.add_run().add_break()
            for seg, b, it in inline_runs(line):
                r = p.add_run(seg)
                r.bold = True if (b or bold) else None
                r.italic = True if it else None
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = docx.shared.RGBColor.from_string(color)
        return p

    def page_break_before(self, p):
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement('w:pageBreakBefore'))

    def bookmark(self, p, name):
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(self.bookmark_id))
        start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(self.bookmark_id))
        self.bookmark_id += 1
        p._p.insert(1 if p._p.pPr is not None else 0, start)
        p._p.append(end)

    def hyperlink(self, p, text, url=None, anchor=None):
        h = OxmlElement('w:hyperlink')
        if anchor:
            h.set(qn('w:anchor'), anchor)
        else:
            rid = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
            h.set(qn('r:id'), rid)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        c = OxmlElement('w:color')
        c.set(qn('w:val'), LINK_COLOR)
        rPr.append(c)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.text = text
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        h.append(r)
        p._p.append(h)

    def field(self, p, instr, placeholder='?'):
        def fld(kind):
            r = OxmlElement('w:r')
            f = OxmlElement('w:fldChar')
            f.set(qn('w:fldCharType'), kind)
            r.append(f)
            return r
        p._p.append(fld('begin'))
        r = OxmlElement('w:r')
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = ' %s ' % instr
        r.append(it)
        p._p.append(r)
        p._p.append(fld('separate'))
        p.add_run(placeholder)
        p._p.append(fld('end'))

    def picture(self, path, width_in, alt):
        p = self.new_par()
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement('w:keepNext'))
        p.paragraph_format.space_after = Pt(5)
        p.add_run().add_picture(path, width=Inches(width_in))
        for docpr in p._p.iter(qn('wp:docPr')):
            docpr.set('descr', alt)
        return p

    def table(self, rows):
        widths = column_widths(rows)
        ncols = len(widths)

        t = self.doc.add_table(rows=len(rows), cols=ncols)
        self.body.remove(t._tbl)
        self.sect.addprevious(t._tbl)
        tblPr = t._tbl.tblPr
        for child in list(tblPr):
            if child.tag == qn('w:tblStyle'):
                tblPr.remove(child)
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        tblPr.append(layout)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(w))

        for ri, (row, data) in enumerate(zip(t.rows, rows)):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if ri == 0:
                trPr.append(OxmlElement('w:tblHeader'))
            fill = TABLE_HEAD_FILL if ri == 0 else (TABLE_BAND_FILL if ri % 2 == 0 else 'FFFFFF')
            for cell, text, w in zip(row.cells, data, widths):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:type'), 'dxa')
                tcW.set(qn('w:w'), str(w))
                va = OxmlElement('w:vAlign')
                va.set(qn('w:val'), 'center')
                tcPr.append(va)
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), fill)
                tcPr.append(shd)
                mar = OxmlElement('w:tcMar')
                for side, v in (('top', 85), ('left', 100), ('bottom', 85), ('right', 100)):
                    e = OxmlElement('w:' + side)
                    e.set(qn('w:w'), str(v))
                    e.set(qn('w:type'), 'dxa')
                    mar.append(e)
                tcPr.append(mar)
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    e = OxmlElement('w:' + side)
                    e.set(qn('w:val'), 'single')
                    e.set(qn('w:sz'), '4')
                    e.set(qn('w:color'), TABLE_BORDER)
                    borders.append(e)
                tcPr.append(borders)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if ri == 0:
                    self.add_text(p, text.replace('**', ''), size=9, color='FFFFFF', bold=True)
                else:
                    self.add_text(p, text, size=9)
        return t


def column_widths(rows):
    """Fixed column widths in twips: every column is at least as wide as
    its longest word (so words never break inside), and the remaining width
    goes to the columns with the most text."""
    ncols = len(rows[0])
    plain = [[re.sub(r'\*+', '', c) for c in r] for r in rows]
    char_tw = [125 if ri == 0 else 105 for ri in range(len(rows))]   # 9 pt, bold header
    pad = 260
    minw = [max(max((len(w) for w in r[c].split()), default=1) * char_tw[ri]
                for ri, r in enumerate(plain)) + pad for c in range(ncols)]
    text = [sum(len(r[c]) for r in plain[1:]) or 1 for c in range(ncols)]
    widths = list(minw)
    spare = TEXT_WIDTH_TWIPS - sum(widths)
    if spare > 0:
        widths = [w + spare * t / sum(text) for w, t in zip(widths, text)]
    else:
        widths = [TEXT_WIDTH_TWIPS * w / sum(widths) for w in widths]
    widths = [round(w) for w in widths]
    widths[-1] = TEXT_WIDTH_TWIPS - sum(widths[:-1])
    return widths


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def fill_cover(b, cover_blocks):
    """Rewrite the template's cover page with the Markdown cover content.
    The cover is everything up to the first page break; its paragraphs keep
    their formatting, only text and logo size change."""
    cover, el = [], b.body[0]
    while el is not None:
        cover.append(el)
        if el.findall('.//' + qn('w:br')) and any(
                br.get(qn('w:type')) == 'page' for br in el.iter(qn('w:br'))):
            break
        el = el.getnext()

    texts = [p for k, p in cover_blocks if k in ('h1', 'para')]
    logo = next(p for k, p in cover_blocks if k == 'img')
    text_pars = [e for e in cover if e.tag == qn('w:p') and ''.join(e.itertext()).strip()]
    if len(text_pars) != len(texts):
        sys.exit('cover mismatch: template has {} text paragraphs, Markdown {}'.format(
            len(text_pars), len(texts)))
    for p_el, text in zip(text_pars, texts):
        p = docx.text.paragraph.Paragraph(p_el, b.doc)
        first = p.runs[0]._r.find(qn('w:rPr')) if p.runs else None
        for r in list(p_el.findall(qn('w:r'))):
            p_el.remove(r)
        b.add_text(p, text)
        if first is not None:
            for r in p_el.findall(qn('w:r')):
                r.insert(0, copy.deepcopy(first))

    if logo['width_px']:
        emu = int(Cm(logo['width_px'] / 96 * 2.54))
        for e in cover:
            for ext in list(e.iter(qn('wp:extent'))) + list(e.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext')):
                ext.set('cx', str(emu))
                ext.set('cy', str(emu))
    return cover


def build():
    md = open(MD_PATH, encoding='utf-8').read()
    blocks = parse_blocks(md)
    b = Builder(DOCX_PATH)

    contents_at = blocks.index(('h2', 'Contents'))
    cover = fill_cover(b, blocks[:contents_at])
    for el in list(b.body):
        if el is not b.sect and el not in cover:
            b.body.remove(el)

    edition = re.search(r'Edition ([\d.]+)\n(.+)', md)
    edition_no, edition_date = edition.group(1), edition.group(2).strip()

    # --- Contents page ---------------------------------------------------
    i = contents_at + 1
    b.add_text(b.new_par('Title'), 'Contents')
    while blocks[i][0] != 'anchor':
        kind, payload = blocks[i]
        if kind == 'bullet':
            m = re.match(r'\[(.+)\]\(#(.+)\)$', payload)
            p = b.new_par('TOCEntry')
            p.paragraph_format.tab_stops.add_tab_stop(
                docx.shared.Twips(9821), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.space_after = Pt(1.75)
            b.hyperlink(p, m.group(1), anchor=m.group(2))
            p.add_run('\t')
            b.field(p, 'PAGEREF %s \\h' % m.group(2))
        elif kind == 'para':
            b.add_text(b.new_par('Small'), payload)
        i += 1
    b.new_par().add_run().add_break(WD_BREAK.PAGE)

    # --- Body ------------------------------------------------------------
    pending_anchor, prev_kind, h1, h2 = None, None, '', ''
    first_h1 = True
    for kind, payload in blocks[i:]:
        if kind == 'anchor':
            pending_anchor = payload
            continue
        if kind == 'h2':                      # Markdown ## -> Heading 1
            h1, h2 = payload, ''
            p = b.add_text(b.new_par('Heading 1'), payload)
            if not first_h1:
                b.page_break_before(p)
            first_h1 = False
        elif kind == 'h3':                    # Markdown ### -> Heading 2
            h2 = payload
            p = b.add_text(b.new_par('Heading 2'), payload)
            if payload in H2_PAGE_BREAKS:
                b.page_break_before(p)
        elif kind == 'img':
            path = os.path.join(MANUAL_DIR, payload['src'])
            if '/gui/' in payload['src']:
                from PIL import Image
                px = Image.open(path).width
                width = min(px / (DETAIL_PX_PER_IN if px < DETAIL_MAX_PX else GUI_PX_PER_IN),
                            TEXT_WIDTH_IN)
            else:
                width = SCHEMATIC_WIDTH_IN
            b.picture(path, width, payload['alt'])
        elif kind == 'para':
            src = re.match(r'\*(Source reference.*)\*$', payload)
            link = re.match(r'\[(.+)\]\((https?://[^)]+)\)$', payload)
            if src:
                b.add_text(b.new_par('SourceNote'), src.group(1))
            elif link:
                b.hyperlink(b.new_par(), link.group(1).replace('\\[', '[').replace('\\]', ']'),
                            url=link.group(2))
            elif prev_kind == 'img' and payload.startswith('Figure '):
                p = b.add_text(b.new_par('Caption'), payload)
                p.paragraph_format.space_after = Pt(6)
            elif payload.startswith('**Figures.**') or _in_small_zone(h1, h2):
                b.add_text(b.new_par('Small'), payload)
            else:
                b.add_text(b.new_par(), payload)
        elif kind == 'num':
            p = b.new_par()
            p.paragraph_format.left_indent = docx.shared.Twips(259)
            p.paragraph_format.first_line_indent = docx.shared.Twips(-259)
            b.add_text(p, '%s.  %s' % payload)
        elif kind == 'bullet':
            p = b.new_par()
            p.paragraph_format.left_indent = docx.shared.Twips(216)
            p.paragraph_format.first_line_indent = docx.shared.Twips(-216)
            b.add_text(p, '•  ' + payload)
        elif kind == 'code':
            b.add_text(b.new_par('Code'), '\n'.join(payload))
        elif kind == 'table':
            b.table(payload)
            # Thin spacer: Word tables have no space-after of their own.
            spacer = b.new_par()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.add_run().font.size = Pt(4)
            spacer.paragraph_format.line_spacing = Pt(6)
        if pending_anchor and kind in ('h2', 'h3'):
            b.bookmark(p, pending_anchor)
            pending_anchor = None
        prev_kind = kind

    # The template's Code style asks for Liberation Mono, which Word on macOS
    # and most Windows machines lack (it then falls back to a serif face).
    b.doc.styles['Code'].font.name = CODE_FONT
    _drop_unused_rels(b.doc)
    _update_footer(b.doc, edition_no, edition_date)
    cp = b.doc.core_properties
    cp.version = edition_no
    cp.modified = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    b.doc.save(DOCX_PATH)
    print('wrote', os.path.relpath(DOCX_PATH, REPO_ROOT))


_SMALL_ZONE = {'on': False}


def _in_small_zone(h1, h2):
    if h1 == SMALL_FROM_H1:
        _SMALL_ZONE['on'] = True
    return _SMALL_ZONE['on'] and h2 not in NORMAL_H2_IN_SMALL_ZONE


def _drop_unused_rels(doc):
    """Remove image and hyperlink relationships left over from the
    template body, so stale media do not stay in the package."""
    xml = doc.element.xml
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.IMAGE, RT.HYPERLINK) and ('"%s"' % rid) not in xml:
            doc.part.drop_rel(rid)


def _update_footer(doc, edition_no, edition_date):
    for section in doc.sections:
        for p in section.footer.paragraphs:
            for t in p._p.iter(qn('w:t')):
                if t.text and t.text.startswith('Edition '):
                    t.text = 'Edition %s  •  %s' % (edition_no, edition_date)


if __name__ == '__main__':
    build()
